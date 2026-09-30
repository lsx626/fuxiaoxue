"""本地全文搜索：CJK 预分词、文档文本抽取与结果过滤。

设计要点（为什么自己分词）：

SQLite FTS5 默认的 ``unicode61`` 分词器不做中文分词——一整段汉字会被
合并成**一个** token。结果是搜「计算机」永远无法命中「计算机体系结构」，
除非在写入前先把汉字串拆开。这里采用查询与索引一致的 **unigram+bigram**
预分词：单字保证单字检索可用，双字保证词组的顺序与精度（「算计」不会
被「计算」误命中）。

提取器只复用桌面预览已有的解析栈（PyMuPDF / python-docx / python-pptx /
openpyxl / odfpy），厚此薄彼地解析常见课件格式；旧二进制 .doc/.ppt/.xls
与图片、音视频不做内容索引（只索引文件名），与预览的降级说明口径一致。
"""
from __future__ import annotations

import os
import re
from html.parser import HTMLParser
from typing import Dict, List, Optional

MAX_INDEX_CHARS = 65536            # 单个文件进入索引的最大字符数
MAX_READ_BYTES = 16 * 1024 * 1024  # 抽取时最多读取的源字节（避免巨型文件拖垮同步）

# CJK 统一表意文字（含扩展 A、兼容表意）与平假名/片假名：这些区间做多字预分词
_CJK_RANGES = (
    (0x3040, 0x30FF),
    (0x3400, 0x4DBF),
    (0x4E00, 0x9FFF),
    (0xF900, 0xFAFF),
)

_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".log",
    ".ini", ".cfg", ".yml", ".yaml", ".py", ".js", ".ts", ".java", ".c",
    ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".sql", ".sh", ".bat",
    ".ps1", ".tex", ".srt", ".vtt", ".toml", ".properties", ".gitignore",
}
_HTML_EXTS = {".html", ".htm", ".xhtml"}
_OFFICE_XML_EXTS = {".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx"}
_ODF_EXTS = {".odt", ".odp", ".ods"}


def _is_cjk(code_point: int) -> bool:
    return any(low <= code_point <= high for low, high in _CJK_RANGES)


# ----------------------------------------------------------------------
# 分词：索引文本与查询文本用同一套规则
# ----------------------------------------------------------------------
def tokenize_for_index(text: str) -> str:
    """把原始文本转成空格分隔的 token 串（unigram+bigram）。

    汉字串「计算机」→ ``计 算 机 计算 算机``。
    """
    if not text:
        return ""
    text = text.lower()
    tokens: List[str] = []
    i = 0
    total = len(text)
    while i < total:
        char = text[i]
        if _is_cjk(ord(char)):
            end = i
            while end < total and _is_cjk(ord(text[end])):
                end += 1
            run = text[i:end]
            tokens.extend(run)                              # 单字
            tokens.extend(run[k:k + 2] for k in range(len(run) - 1))  # 双字
            i = end
        elif char.isalnum():
            end = i
            while end < total and text[end].isalnum() and not _is_cjk(ord(text[end])):
                end += 1
            tokens.append(text[i:end])
            i = end
        else:
            i += 1
    return " ".join(tokens)


def build_match_query(query: str) -> Optional[str]:
    """把用户输入转成 FTS5 MATCH 表达式；空查询返回 None。

    所有 token 显式加引号，避免用户输入中的特殊字符破坏 FTS5 语法；
    最后一个 ASCII 词追加 ``*`` 前缀匹配（输入 ``lec`` 能命中 ``lecture5``）。
    """
    if not query:
        return None
    query = query.lower()
    terms: List[str] = []
    i = 0
    total = len(query)
    while i < total:
        char = query[i]
        if _is_cjk(ord(char)):
            end = i
            while end < total and _is_cjk(ord(query[end])):
                end += 1
            run = query[i:end]
            if len(run) == 1:
                terms.append(run)
            else:
                terms.extend(run)
                terms.extend(run[k:k + 2] for k in range(len(run) - 1))
            i = end
        elif char.isalnum():
            end = i
            while end < total and query[end].isalnum() and not _is_cjk(ord(query[end])):
                end += 1
            terms.append(query[i:end])
            i = end
        else:
            i += 1
    if not terms:
        return None
    terms = terms[:64]
    parts = []
    last = len(terms) - 1
    for index, term in enumerate(terms):
        if index == last and term.isascii() and term.isalnum() and len(term) >= 2:
            parts.append(f'"{term}"*')
        else:
            parts.append(f'"{term}"')
    return " ".join(parts)


# ----------------------------------------------------------------------
# 文档文本抽取
# ----------------------------------------------------------------------
def extract_text(path: str) -> Optional[str]:
    """按扩展名抽取纯文本；失败或格式不支持返回 None（调用方只索引文件名）。"""
    if not path or not os.path.isfile(path):
        return None
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in _TEXT_EXTS:
            return _extract_plain(path)
        if ext in _HTML_EXTS:
            return _extract_html(path)
        if ext == ".pdf":
            return _extract_pdf(path)
        if ext in _OFFICE_XML_EXTS:
            kind = _OFFICE_XML_EXTS[ext]
            return _extract_office_xml(path, kind)
        if ext in _ODF_EXTS:
            return _extract_odf(path)
    except Exception:  # pylint: disable=broad-except
        # 抽取失败不能影响同步：返回 None，索引只写文件名
        return None
    return None


def _decode_text(raw: bytes) -> str:
    """带回退地解码文本：UTF-8-sig → GB18030 → GBK → ISO-8859-1 兜底。

    桌面预览固定 ``utf-8-sig`` 导致 GBK 记事本文件全屏乱码（Android 端
    v1.0.12 已修）；索引层不能再犯同一个错。
    """
    for encoding in ("utf-8-sig", "gb18030", "gbk"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("latin-1", errors="replace")


def _read_capped(path: str) -> bytes:
    with open(path, "rb") as handle:
        return handle.read(MAX_READ_BYTES)


def _truncate(text: str) -> str:
    return text[:MAX_INDEX_CHARS] if len(text) > MAX_INDEX_CHARS else text


def _extract_plain(path: str) -> Optional[str]:
    return _truncate(_decode_text(_read_capped(path))).strip() or None


class _HtmlTextExtractor(HTMLParser):
    """HTML → 纯文本：跳过 script/style，转换实体，折叠空白。"""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._parts: List[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs) -> None:  # noqa: N802 - html.parser API
        if tag in ("script", "style"):
            self._skip_depth += 1
        self._parts.append(" ")

    def handle_endtag(self, tag) -> None:  # noqa: N802 - html.parser API
        if tag in ("script", "style") and self._skip_depth:
            self._skip_depth -= 1
        self._parts.append(" ")

    def handle_data(self, data) -> None:  # noqa: N802 - html.parser API
        if not self._skip_depth:
            self._parts.append(data)

    def text(self) -> str:
        return re.sub(r"\s+", " ", "".join(self._parts)).strip()


def _extract_html(path: str) -> Optional[str]:
    parser = _HtmlTextExtractor()
    parser.feed(_decode_text(_read_capped(path)))
    parser.close()
    return _truncate(parser.text()) or None


def _extract_pdf(path: str) -> Optional[str]:
    import fitz  # noqa: PLC0415 - 懒加载，缺 PyMuPDF 时不影响其它格式
    parts: List[str] = []
    length = 0
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text()
            parts.append(text)
            length += len(text)
            if length >= MAX_INDEX_CHARS:
                break
    return _truncate("\n".join(parts)).strip() or None


def _extract_office_xml(path: str, kind: str) -> Optional[str]:
    parts: List[str] = []

    def append(text: str) -> None:
        if text:
            parts.append(text)

    if kind == "docx":
        import docx  # noqa: PLC0415
        document = docx.Document(path)
        for paragraph in document.paragraphs:
            append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    append(cell.text)
    elif kind == "pptx":
        import pptx  # noqa: PLC0415
        presentation = pptx.Presentation(path)
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    append(shape.text_frame.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            append(cell.text)
    else:  # xlsx
        import openpyxl  # noqa: PLC0415
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows(values_only=True):
                    for value in row:
                        if value is not None:
                            append(str(value))
        finally:
            workbook.close()
    return _truncate(" ".join(parts)).strip() or None


def _extract_odf(path: str) -> Optional[str]:
    import odf.opendocument  # noqa: PLC0415
    import odf.teletype  # noqa: PLC0415
    document = odf.opendocument.load(path)
    return _truncate(odf.teletype.extractText(document)).strip() or None


# ----------------------------------------------------------------------
# 文件列表即时过滤（与 FTS 无关的纯函数，供文件表过滤框使用）
# ----------------------------------------------------------------------
def filter_file_records(files: List[Dict], query: str) -> List[Dict]:
    """按空格分隔的关键字过滤文件记录（匹配文件名/来源/上下文/目录）。"""
    terms = [term.lower() for term in (query or "").split() if term]
    if not terms:
        return list(files)
    results = []
    for file_info in files:
        haystack = " ".join(
            str(file_info.get(key) or "")
            for key in ("filename", "display_name", "source", "context", "folder_path")
        ).lower()
        if all(term in haystack for term in terms):
            results.append(file_info)
    return results
