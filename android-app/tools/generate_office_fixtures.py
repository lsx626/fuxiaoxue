"""Generate synthetic Android legacy Office fixtures on Windows with Office.

Run using the desktop project's packaging environment, then adb push the three
files into the dedicated test app's external files/fixtures directory. No real
documents are opened, macros are disabled, and only owned COM instances close.
"""
from __future__ import annotations

import argparse
from pathlib import Path


def generate(output: Path) -> None:
    from PIL import Image, ImageDraw
    from docx import Document
    from docx.shared import Inches, RGBColor
    from pptx import Presentation
    from pptx.util import Inches as PptInches
    import win32com.client

    output.mkdir(parents=True, exist_ok=True)
    picture = output / "synthetic.png"
    image = Image.new("RGB", (320, 180), "#4f46e5")
    ImageDraw.Draw(image).rectangle((40, 40, 280, 140), fill="#e0e3ff")
    image.save(picture)

    document = Document()
    document.add_heading("Synthetic Office fixture", 0)
    run = document.add_paragraph().add_run("Formatted text for Android preview testing")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 40, 40)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "95"
    document.add_picture(str(picture), width=Inches(2.5))
    modern = output / "modern.docx"
    document.save(modern)

    word = win32com.client.DispatchEx("Word.Application")
    doc = None
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        doc = word.Documents.Open(str(modern), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        doc.SaveAs2(str(output / "legacy.doc"), FileFormat=0, AddToRecentFiles=False)
    finally:
        if doc is not None:
            doc.Close(SaveChanges=0)
        word.Quit(SaveChanges=0)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(PptInches(.5), PptInches(.3), PptInches(8), PptInches(1)).text = "Synthetic slide one"
    cells = slide.shapes.add_table(2, 2, PptInches(.5), PptInches(1.5), PptInches(4), PptInches(1.5)).table
    cells.cell(0, 0).text = "Name"
    cells.cell(0, 1).text = "Score"
    cells.cell(1, 0).text = "Alice"
    cells.cell(1, 1).text = "95"
    slide.shapes.add_picture(str(picture), PptInches(5), PptInches(1.5), width=PptInches(3))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(PptInches(.5), PptInches(.5), PptInches(8), PptInches(1)).text = "Synthetic slide two"
    pptx = output / "modern.pptx"
    presentation.save(pptx)
    powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
    slides = None
    try:
        powerpoint.AutomationSecurity = 3
        powerpoint.DisplayAlerts = 1
        slides = powerpoint.Presentations.Open(str(pptx), ReadOnly=True, Untitled=True, WithWindow=False)
        slides.SaveAs(str(output / "legacy.ppt"), 1)
    finally:
        if slides is not None:
            slides.Close()
        powerpoint.Quit()
    for name in ("legacy.doc", "modern.docx", "legacy.ppt"):
        target = output / name
        if not target.is_file() or not target.stat().st_size:
            raise RuntimeError(f"Fixture missing: {name}")
        print(f"Generated {name} ({target.stat().st_size} bytes)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    generate(parser.parse_args().output.resolve())
